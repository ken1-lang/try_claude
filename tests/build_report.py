# テスト結果(tests/output/results.json)から「テスト仕様書兼結果表」の Excel を生成する。
# 使い方: python3 tests/build_report.py [出力先.xlsx]
import json, sys, datetime, subprocess, os
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
cases = [(c['category'], c['item'], c['steps'], c['expected']) for c in json.load(open(os.path.join(HERE, 'cases.json'), encoding='utf-8'))]
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'calculator_test_spec.xlsx')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = "Meiryo"
data = json.load(open(os.path.join(HERE, 'output', 'results.json'), encoding='utf-8'))
res = {r['no']: r for r in data['results']}
meta = data['meta']
commit = subprocess.check_output(['git', '-C', ROOT, 'rev-parse', '--short', 'HEAD']).decode().strip()
branch = subprocess.check_output(['git', '-C', ROOT, 'branch', '--show-current']).decode().strip()
today = datetime.date.today()

wb = Workbook()
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill("solid", fgColor="1F3864")

ws = wb.active; ws.title = "テスト項目"
heads = ["No.", "分類", "テスト項目", "操作手順", "期待結果", "実際の結果", "結果", "実施日", "実施者", "備考"]
widths = [7, 16, 28, 40, 50, 46, 10, 12, 18, 40]
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = hfill; c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 24
for n, (cat, item, step, exp) in enumerate(cases, 1):
    no = f"T{n:03d}"; r = n + 1; x = res[no]; assert x['item'] == item
    done = x['result'] in ('OK', 'NG')
    vals = [no, cat, item, step, exp, x['actual'] if done else None, x['result'],
            today if done else None, "Claude(自動テスト)" if done else None, x['note'] or None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=10); c.border = border
        c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if i in (1, 7, 8) else "left")
    ws.cell(row=r, column=8).number_format = "yyyy/mm/dd"
last = len(cases) + 1
dv = DataValidation(type="list", formula1='"OK,NG,未実施"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"G2:G{last}")
rng = f"G2:G{last}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(name=F, color="006100")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NG"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=F, color="9C0006")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"未実施"'], fill=PatternFill("solid", bgColor="FFEB9C"), font=Font(name=F, color="9C5700")))
ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:J{last}"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "1:1"

ov = wb.create_sheet("概要", 0)
ov.column_dimensions["A"].width = 22; ov.column_dimensions["B"].width = 70
def put(r, a, b):
    x = ov.cell(row=r, column=1, value=a); y = ov.cell(row=r, column=2, value=b)
    x.font = Font(name=F, size=10, bold=True); y.font = Font(name=F, size=10)
    x.fill = PatternFill("solid", fgColor="D9E1F2")
    for c in (x, y): c.border = border; c.alignment = Alignment(vertical="top", wrap_text=True)
ov.cell(row=1, column=1, value="電卓アプリ テスト仕様書兼結果表").font = Font(name=F, size=14, bold=True)
put(3, "対象", "電卓アプリ(index.html)")
put(4, "対象バージョン", f"ブランチ {branch}(コミット {commit} 時点の index.html)")
put(5, "公開URL", "https://ken1-lang.github.io/try_claude/")
put(6, "対象機能", "四則演算、クリア、1文字削除、パーセント、履歴、キーボード操作")
put(7, "テスト環境", f"{meta['browser']}(Playwright による自動操作)。画面サイズ 800×900(スマホ幅確認のみ 360×640)。対象: {meta['url']}")
put(8, "実施日", today)
ov.cell(row=8, column=2).number_format = "yyyy/mm/dd"; ov.cell(row=8, column=2).alignment = Alignment(horizontal="left")
put(9, "実施方法", "各項目の操作手順をブラウザ上のボタン/キーボード操作で自動実行し、画面の表示を期待結果と照合した。各テストは新しいブラウザ環境(履歴なし)から開始")
put(10, "未実施の項目", "Edge、iPhone の Safari、Android の Chrome、公開URL(テスト環境から接続不可)は、実機または当該環境が必要なため未実施。手動での確認が必要")
put(12, "総項目数", f"=COUNTA(テスト項目!A2:A{last})")
put(13, "OK", f'=COUNTIF(テスト項目!G2:G{last},"OK")')
put(14, "NG", f'=COUNTIF(テスト項目!G2:G{last},"NG")')
put(15, "未実施", "=B12-B13-B14")
put(16, "実施率", "=IF(B12=0,0,(B13+B14)/B12)")
put(17, "合格率(OK/実施済)", "=IF(B13+B14=0,0,B13/(B13+B14))")
for r in (16, 17): ov.cell(row=r, column=2).number_format = "0.0%"
for r in range(12, 18): ov.cell(row=r, column=2).alignment = Alignment(horizontal="left")
wb.save(OUT)
print("saved", OUT)
